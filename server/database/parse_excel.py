# server/database/parse_excel.py
import openpyxl
import json
import re
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

def clean_val(val):
    if val is None:
        return ""
    val_str = str(val).strip()
    if val_str.startswith('='):
        # Extract hyperlink URL if present
        hyperlink_match = re.search(r'HYPERLINK\("([^"]+)"', val_str, re.IGNORECASE)
        if hyperlink_match:
            return hyperlink_match.group(1)
        
        # Check for quoted strings in formulas (often containing default/computed strings or fallback URLs)
        str_matches = re.findall(r'"([^"]+)"', val_str)
        if str_matches:
            # Look for URLs first
            for s in reversed(str_matches):
                if s.startswith('http') or 'x.com' in s or 'instagram.com' in s or 'youtube.com' in s or 'facebook.com' in s:
                    return s
            # Fallback to the last match
            return str_matches[-1]
        return ""
    return val_str

def parse_brands():
    path = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../../Kunal_brands.xlsx'))
    print(f"Parsing brands from: {path}")
    wb = openpyxl.load_workbook(path, read_only=True)
    sheet = wb['Sheet3']
    
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]
    
    brands_data = []
    # Skip header
    for r_idx, row in enumerate(rows[1:], start=2):
        # If the row is empty or name is missing, skip
        if not row or len(row) < 2 or not row[1]:
            continue
            
        brand_name = clean_val(row[1])
        if not brand_name:
            continue
            
        brands_data.append({
            "brand_name": brand_name,
            "founded_year": clean_val(row[2]) if len(row) > 2 else "",
            "category": clean_val(row[3]) if len(row) > 3 else "",
            "brand_focus": clean_val(row[4]) if len(row) > 4 else "",
            "founder_names": clean_val(row[5]) if len(row) > 5 else "",
            "revenue": clean_val(row[6]) if len(row) > 6 else "",
            "revenue_year": clean_val(row[7]) if len(row) > 7 else "",
            "last_funding_amount": clean_val(row[8]) if len(row) > 8 else "",
            "last_funding_data": clean_val(row[9]) if len(row) > 9 else "",
            "last_funding_date": clean_val(row[10]) if len(row) > 10 else "",
            "headquarter": clean_val(row[11]) if len(row) > 11 else "",
            "main_geography_outreach": clean_val(row[12]) if len(row) > 12 else "",
            "linkedin": clean_val(row[13]) if len(row) > 13 else "",
            "how_many_employees": clean_val(row[14]) if len(row) > 14 else "",
            "marketing_head": clean_val(row[15]) if len(row) > 15 else "",
            "marketing_mail_id": clean_val(row[16]) if len(row) > 16 else "",
            "sales_head": clean_val(row[17]) if len(row) > 17 else "",
            "sales_head_mail": clean_val(row[18]) if len(row) > 18 else "",
            "content_marketing_head": clean_val(row[19]) if len(row) > 19 else "",
            "content_marketing_head_mail_id": clean_val(row[20]) if len(row) > 20 else "",
            "company_phone": clean_val(row[21]) if len(row) > 21 else "",
            "company_url": clean_val(row[22]) if len(row) > 22 else "",
            "facebook": clean_val(row[23]) if len(row) > 23 else "",
            "instagram": clean_val(row[24]) if len(row) > 24 else "",
            "youtube": clean_val(row[25]) if len(row) > 25 else "",
            "twitter": clean_val(row[26]) if len(row) > 26 else "",
            "main_influencer_platform": clean_val(row[27]) if len(row) > 27 else "",
            "web_traffic": clean_val(row[28]) if len(row) > 28 else ""
        })
        
    out_path = os.path.join(os.path.dirname(__file__), 'brands.json')
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(brands_data, f, ensure_ascii=False, indent=2)
    print(f"Successfully wrote {len(brands_data)} brands to {out_path}")

def parse_influencers():
    path = os.path.abspath(os.path.join(os.path.dirname(__file__), '../../../_Outreach Strategy.xlsx'))
    print(f"Parsing influencers from: {path}")
    wb = openpyxl.load_workbook(path, data_only=True) # data_only works fine here because we just have plain fields mostly
    sheet = wb['Influencer Outreach']
    
    rows = list(sheet.iter_rows(values_only=True))
    headers = rows[0]
    
    influencers_data = []
    for r_idx, row in enumerate(rows[1:], start=2):
        if not row or len(row) < 1 or not row[0]:
            continue
            
        name = str(row[0]).strip()
        if not name or name == "Influencer Name":
            continue
            
        send_date_val = row[7]
        send_date_str = None
        if send_date_val:
            if hasattr(send_date_val, 'strftime'):
                send_date_str = send_date_val.strftime('%Y-%m-%d')
            else:
                send_date_str = str(send_date_val).split(' ')[0] # try to extract date part
                
        influencers_data.append({
            "influencer_name": name,
            "lead_by": str(row[1]).strip() if len(row) > 1 and row[1] is not None else "",
            "content_why_this_person": str(row[2]).strip() if len(row) > 2 and row[2] is not None else "",
            "instagram_url": str(row[3]).strip() if len(row) > 3 and row[3] is not None else "",
            "followers": str(row[4]).strip() if len(row) > 4 and row[4] is not None else "",
            "script": str(row[5]).strip() if len(row) > 5 and row[5] is not None else "",
            "comment_average": str(row[6]).strip() if len(row) > 6 and row[6] is not None else "",
            "send_date": send_date_str
        })
        
    out_path = os.path.join(os.path.dirname(__file__), 'influencers.json')
    with open(out_path, 'w', encoding='utf-8') as f:
        json.dump(influencers_data, f, ensure_ascii=False, indent=2)
    print(f"Successfully wrote {len(influencers_data)} influencers to {out_path}")

if __name__ == '__main__':
    parse_brands()
    parse_influencers()
